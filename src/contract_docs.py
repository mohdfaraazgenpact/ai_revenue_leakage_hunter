"""Build realistic synthetic contract documents (clause text + PDF/DOCX rendering).

Escalation obligations are written in many different ways - explicit dates, relative
"anniversary of the Effective Date" wording, numbers spelled out, clauses buried in the
Miscellaneous section, terms inside Schedule A, CPI indexation with caps/floors. This is
what makes contract reading hard for rules and is the showcase for LLM extraction.
Distractor percentages (late-payment interest, liability caps) are included on purpose.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

SUPPLIER = "Orion Global Solutions Inc."

_WORDS = {
    "1": "one",
    "1.5": "one and one-half",
    "2": "two",
    "2.5": "two and one-half",
    "3": "three",
    "3.5": "three and one-half",
    "4": "four",
    "4.5": "four and one-half",
    "5": "five",
    "6": "six",
    "7.5": "seven and one-half",
    "10": "ten",
}


def pct_str(p: Decimal) -> str:
    s = format(p.normalize(), "f")
    return s


def pct_words(p: Decimal) -> str:
    return _WORDS[pct_str(p)]


def fmt_date(d: date, style: int) -> str:
    if style == 0:
        return f"{d.day} {d:%B %Y}"
    return f"{d:%B} {d.day}, {d.year}"


@dataclass
class Clause:
    clause_id: str
    text: str


@dataclass
class Section:
    heading: str
    clauses: list[Clause]


@dataclass
class ContractDoc:
    contract_id: str
    title: str
    preamble: str
    sections: list[Section]
    schedule: Section
    attachments: list[Section] = field(default_factory=list)  # amendments, side letters
    # field name -> (clause_id, quote) for gold citations
    citations: dict[str, tuple[str, str]] = field(default_factory=dict)

    def all_clauses(self):
        for s in [*self.sections, self.schedule, *self.attachments]:
            yield from s.clauses


def escalation_clause(spec, style: int) -> tuple[str, str]:
    """Return (placement, text). placement: '4.2' | '13.4' | 'schedule'."""
    esc = spec["escalation"]
    d1 = esc.get("first_escalation_date")
    D1 = fmt_date(d1, style) if d1 else ""
    v = spec["wording_variant"]
    if esc["type"] == "none":
        return "4.2", (
            "The Unit Prices are fixed for the Term and are not subject to any escalation or indexation."
            if v % 2 == 0
            else "Unit Prices shall remain firm and fixed for the duration of this Agreement."
        )
    if esc["type"] == "cpi":
        lo, hi = esc["floor_pct"], esc["cap_pct"]
        if v % 2 == 0:
            return "4.2", (
                "On each anniversary of the Effective Date, the Unit Prices shall be adjusted by the "
                "percentage change in the Consumer Price Index for All Urban Consumers (CPI-U) for the "
                f"preceding calendar year, provided that no such adjustment shall be less than {pct_str(lo)}% "
                f"or greater than {pct_str(hi)}%."
            )
        return "4.2", (
            f"Commencing {D1} and annually thereafter, Unit Prices shall be revised in line with the annual "
            "CPI-U inflation rate published for the prior calendar year, subject to a minimum increase of "
            f"{pct_words(lo)} percent ({pct_str(lo)}%) and a maximum increase of {pct_words(hi)} percent "
            f"({pct_str(hi)}%)."
        )
    p = esc["rate_pct"]
    P, W = pct_str(p), pct_words(p)
    month_day = f"{d1:%B} {d1.day}" if d1 else ""
    texts = {
        1: (
            "4.2",
            f"On {D1} and on each anniversary of that date thereafter, the Unit Prices shall increase by {P}%.",
        ),
        2: (
            "4.2",
            f"Unit Prices are subject to an annual increase of {W} percent ({P}%) effective each {month_day}, commencing {D1}.",
        ),
        3: (
            "4.2",
            "Commencing on the first anniversary of the Effective Date and on each anniversary thereafter, "
            f"all Unit Prices in Schedule A shall be adjusted upward by {W} percent ({P}%), compounding.",
        ),
        4: (
            "4.2",
            "Following the end of the first Contract Year, the Unit Prices shall increase at the start of "
            f"each subsequent Contract Year by {W} percent.",
        ),
        5: (
            "4.2",
            f"The parties agree to an annual price uplift of {P}% which shall take effect on {D1} and annually thereafter.",
        ),
        6: (
            "13.4",
            f"Notwithstanding Section 4.2, the Unit Prices in Schedule A will escalate at {W} percent ({P}%) "
            f"per annum, with the first such escalation on {D1} and subsequent escalations on each anniversary thereof.",
        ),
        7: (
            "schedule",
            f"Annual Adjustment: the Unit Prices above increase by {P}% compounding, first applied on {D1} "
            "and on each anniversary thereafter.",
        ),
        8: (
            "4.2",
            f"Pricing shall remain firm through {fmt_date(d1 - timedelta(days=1), style)}. Thereafter, Unit Prices "
            f"shall increase by {P}% on {D1} and on each anniversary of that date.",
        ),
    }
    return texts[v]


def build_contract(spec) -> ContractDoc:
    style = spec["date_style"]
    eff, end = spec["effective_date"], spec["end_date"]
    cust = spec["customer_name"]
    cites: dict[str, tuple[str, str]] = {}
    placement, esc_text = escalation_clause(spec, style)

    pricing_42 = {
        "4.2": esc_text,
        "13.4": "Unit Prices shall remain firm except as otherwise expressly provided in this Agreement.",
        "schedule": "Unit Prices are subject to the annual adjustment specified in Schedule A.",
    }[placement]

    eff_txt = f'"Effective Date" means {fmt_date(eff, style)}.'
    term_txt = (
        "This Agreement commences on the Effective Date and continues until "
        f"{fmt_date(end, style)}, unless terminated earlier in accordance with Section 10."
    )
    sections = [
        Section(
            "1. DEFINITIONS",
            [
                Clause(
                    "1.1",
                    '"Agreement" means this Master Supply Agreement, including its Schedules, amendments and side letters.',
                ),
                Clause("1.2", eff_txt),
                Clause(
                    "1.3",
                    '"Contract Year" means each consecutive twelve (12) month period commencing on the Effective Date or an anniversary of it.',
                ),
                Clause(
                    "1.4",
                    '"Products" means the goods and services listed in Schedule A, and "Unit Price" means the price per unit stated there.',
                ),
            ],
        ),
        Section("2. TERM", [Clause("2.1", term_txt)]),
        Section(
            "3. ORDERS AND DELIVERY",
            [
                Clause(
                    "3.1",
                    "Customer may place purchase orders for Products from time to time. Each accepted purchase order is governed exclusively by this Agreement.",
                ),
                Clause(
                    "3.2",
                    "Supplier shall deliver Products in accordance with the delivery schedule in the applicable purchase order. Title and risk pass on delivery.",
                ),
            ],
        ),
        Section(
            "4. PRICING",
            [
                Clause(
                    "4.1",
                    "Customer shall pay the Unit Prices set out in Schedule A. All prices are stated in United States Dollars (USD) and are exclusive of applicable taxes.",
                ),
                Clause("4.2", pricing_42),
            ],
        ),
        Section(
            "5. INVOICING AND PAYMENT",
            [
                Clause(
                    "5.1",
                    "Supplier shall invoice Customer monthly in arrears for Products delivered during the preceding month.",
                ),
                Clause("5.2", "Payment is due net forty-five (45) days from the invoice date."),
                Clause(
                    "5.3",
                    "Overdue amounts accrue interest at 1% per month, or the maximum rate permitted by law if lower.",
                ),
            ],
        ),
        Section(
            "6. WARRANTY",
            [
                Clause(
                    "6.1",
                    "Supplier warrants that Products will conform to their specifications for twelve (12) months from delivery.",
                )
            ],
        ),
        Section(
            "7. LIMITATION OF LIABILITY",
            [
                Clause(
                    "7.1",
                    "Each party's aggregate liability shall not exceed one hundred percent (100%) of the fees paid in the twelve (12) months preceding the claim.",
                )
            ],
        ),
        Section(
            "8. CONFIDENTIALITY",
            [
                Clause(
                    "8.1",
                    "Each party shall keep the other party's Confidential Information confidential and use it only to perform this Agreement.",
                )
            ],
        ),
        Section(
            "9. INSURANCE",
            [
                Clause(
                    "9.1",
                    "Supplier shall maintain commercial general liability insurance of not less than USD 5,000,000 per occurrence.",
                )
            ],
        ),
        Section(
            "10. TERMINATION",
            [
                Clause(
                    "10.1",
                    "Either party may terminate this Agreement for material breach not cured within thirty (30) days of written notice.",
                ),
                Clause(
                    "10.2",
                    "Termination does not affect accrued rights, including the right to payment for Products delivered.",
                ),
            ],
        ),
        Section(
            "11. FORCE MAJEURE",
            [
                Clause(
                    "11.1",
                    "Neither party is liable for delay caused by events beyond its reasonable control.",
                )
            ],
        ),
        Section(
            "12. GOVERNING LAW",
            [Clause("12.1", "This Agreement is governed by the laws of the State of New York.")],
        ),
        Section(
            "13. MISCELLANEOUS",
            [
                Clause(
                    "13.1",
                    "This Agreement is the entire agreement of the parties regarding its subject matter.",
                ),
                Clause(
                    "13.2",
                    "Notices must be in writing and delivered to the addresses stated above.",
                ),
                Clause(
                    "13.3",
                    "This Agreement may be amended only by a written instrument signed by both parties.",
                ),
            ],
        ),
    ]
    if placement == "13.4":
        sections[-1].clauses.append(Clause("13.4", esc_text))

    cites["customer_name"] = ("preamble", cust)
    cites["effective_date"] = ("1.2", eff_txt)
    cites["end_date"] = ("2.1", term_txt)
    if spec["escalation"]["type"] != "none":
        cites["escalation"] = ({"schedule": "A.99"}.get(placement, placement), esc_text)
    else:
        cites["escalation"] = ("4.2", esc_text)

    sched = []
    for i, s in enumerate(spec["skus"], start=1):
        line = f"{s['sku']} - {s['description']}: USD {s['unit_price']:,.2f} per unit."
        sched.append(Clause(f"A.{i}", line))
        cites[f"price:{s['sku']}"] = (f"A.{i}", line)
    if placement == "schedule":
        sched.append(Clause("A.99", esc_text))
    schedule = Section("SCHEDULE A - PRODUCTS AND UNIT PRICES", sched)

    attachments = []
    for k, am in enumerate(spec.get("amendments", []), start=1):
        sku = next(s for s in spec["skus"] if s["sku"] == am["sku"])
        text = (
            f"With effect from {fmt_date(am['effective_date'], style)}, the Unit Price for {am['sku']} "
            f"({sku['description']}) in Schedule A shall be USD {am['new_unit_price']:,.2f} per unit. "
            "Price adjustments under the Agreement continue to apply to the amended Unit Price on subsequent anniversaries."
        )
        cid = f"AM{k}.1"
        attachments.append(
            Section(
                f"AMENDMENT NO. {k} (dated {fmt_date(am['signed_date'], style)})",
                [
                    Clause(cid, text),
                    Clause(
                        f"AM{k}.2",
                        "All other terms of the Agreement remain unchanged and in full force and effect.",
                    ),
                ],
            )
        )
        cites[f"amendment:{k}"] = (cid, text)
    for k, dsc in enumerate(spec.get("approved_discounts", []), start=1):
        sku = next(s for s in spec["skus"] if s["sku"] == dsc["sku"])
        text = (
            f"For the period from {fmt_date(dsc['start_date'], style)} to {fmt_date(dsc['end_date'], style)} (inclusive), "
            f"Supplier grants Customer a promotional discount of {pct_words(dsc['discount_pct'])} percent "
            f"({pct_str(dsc['discount_pct'])}%) on the Unit Price of {dsc['sku']} ({sku['description']}). "
            "The discount applies only during the stated period."
        )
        cid = f"SL{k}.1"
        attachments.append(Section(f"SIDE LETTER NO. {k}", [Clause(cid, text)]))
        cites[f"discount:{k}"] = (cid, text)

    preamble = (
        f"This Master Supply Agreement (Contract No. {spec['contract_id']}) is entered into between "
        f'{SUPPLIER} ("Supplier") and {cust} ("Customer").'
    )
    return ContractDoc(
        contract_id=spec["contract_id"],
        title="MASTER SUPPLY AGREEMENT",
        preamble=preamble,
        sections=sections,
        schedule=schedule,
        attachments=attachments,
        citations=cites,
    )


# ---------------------------------------------------------------- rendering
def _latin1(s: str) -> str:
    return s.encode("latin-1", "replace").decode("latin-1")


def render_pdf(doc: ContractDoc, path: Path) -> None:
    from fpdf import FPDF

    pdf = FPDF(format="Letter")
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.set_margins(22, 20, 22)
    pdf.add_page()
    pdf.set_font("Times", "B", 15)
    pdf.cell(0, 10, doc.title, align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Times", "", 9)
    pdf.cell(
        0,
        5,
        "SYNTHETIC DOCUMENT - FOR POC DEMONSTRATION ONLY",
        align="C",
        new_x="LMARGIN",
        new_y="NEXT",
    )
    pdf.ln(4)
    pdf.set_font("Times", "", 11)
    pdf.multi_cell(0, 5.5, _latin1(doc.preamble), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    def write_section(sec: Section, new_page: bool = False):
        if new_page:
            pdf.add_page()
        pdf.set_font("Times", "B", 11.5)
        pdf.multi_cell(0, 6, _latin1(sec.heading), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Times", "", 11)
        for c in sec.clauses:
            pdf.multi_cell(
                0, 5.5, _latin1(f"{c.clause_id}  {c.text}"), new_x="LMARGIN", new_y="NEXT"
            )
            pdf.ln(1.5)
        pdf.ln(2)

    for s in doc.sections:
        write_section(s)
    pdf.ln(4)
    pdf.multi_cell(
        0,
        5.5,
        f"Signed for and on behalf of {SUPPLIER} and the Customer by their authorised signatories.",
    )
    write_section(doc.schedule, new_page=True)
    for a in doc.attachments:
        write_section(a, new_page=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))


def render_docx(doc: ContractDoc, path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading(doc.title, level=1)
    d.add_paragraph("SYNTHETIC DOCUMENT - FOR POC DEMONSTRATION ONLY")
    d.add_paragraph(doc.preamble)
    for s in [*doc.sections, doc.schedule, *doc.attachments]:
        d.add_heading(s.heading, level=2)
        for c in s.clauses:
            d.add_paragraph(f"{c.clause_id}  {c.text}")
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))
